from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_LINE_SPACING
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont

OUT = Path('/Users/aiot/Documents/workspace/machine learning/ml-lab/output/doc')
OUT.mkdir(parents=True, exist_ok=True)
path = OUT / 'ML_Lab_Bao_cao_Logistic_Regression.docx'
FIG = OUT.parent.parent / 'tmp' / 'sigmoid_curve.png'

GREEN = '173B34'
MOSS = '6E8F68'
ORANGE = 'D97845'
INK = '1F2925'
MUTED = '5E6B64'
CREAM = 'F7F3EA'
PALE = 'E7EFE3'
CORAL = 'F5E4DC'

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_border(cell, **kwargs):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        if edge in kwargs:
            tag = 'w:{}'.format(edge); element = borders.find(qn(tag))
            if element is None: element = OxmlElement(tag); borders.append(element)
            for key in ['val','sz','space','color']:
                if key in kwargs[edge]: element.set(qn('w:'+key), str(kwargs[edge][key]))

def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m, v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn('w:'+m))
        if node is None: node = OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'),'dxa')

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr(); tblHeader = OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'),'true'); trPr.append(tblHeader)

def add_run(p, text, bold=False, color=None, size=None, italic=False):
    r = p.add_run(text); r.bold=bold; r.italic=italic
    if color: r.font.color.rgb = RGBColor.from_string(color)
    if size: r.font.size = Pt(size)
    return r

def add_heading(doc, text, level=1, number=None):
    p = doc.add_paragraph(style=f'Heading {level}')
    if number: add_run(p, f'{number}  ', bold=True, color=INK)
    add_run(p, text, bold=True, color=INK)
    return p

def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph(style='Body Text')
    if bold_lead and text.startswith(bold_lead):
        add_run(p, bold_lead, bold=True, color=GREEN); add_run(p, text[len(bold_lead):])
    else: add_run(p, text)
    return p

def add_callout(doc, title, text, fill=PALE, accent=MOSS):
    p=doc.add_paragraph(style='Body Text'); p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(9)
    add_run(p, title + ': ', bold=True, color=INK); add_run(p, text, color=INK)

def add_bullets(doc, items):
    for item in items:
        p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(4); add_run(p,item)

def make_sigmoid_figure():
    w,h=1200,430; img=Image.new('RGB',(w,h),'#F7F3EA'); d=ImageDraw.Draw(img)
    left,right,top,bottom=110,1120,45,355
    d.line((left,bottom,right,bottom), fill='#6E8F68', width=3); d.line((left,top,left,bottom), fill='#6E8F68', width=3)
    pts=[]
    import math
    for i in range(241):
        x=-6+i*12/240; y=1/(1+math.exp(-x)); px=left+(x+6)/12*(right-left); py=bottom-y*(bottom-top); pts.append((px,py))
    d.line(pts, fill='#D97845', width=6, joint='curve')
    font_path='/System/Library/Fonts/Supplemental/Arial.ttf'
    try: font=ImageFont.truetype(font_path,24); small=ImageFont.truetype(font_path,20)
    except: font=small=None
    d.text((right-80,bottom+12),'z',fill='#1F2925',font=font); d.text((left-55,top-5),'p',fill='#1F2925',font=font)
    d.text((right-35,bottom+12),'→',fill='#1F2925',font=small); d.text((left-55,top-28),'1',fill='#1F2925',font=small); d.text((left-55,bottom-12),'0',fill='#1F2925',font=small)
    d.line((left+(right-left)/2,top,left+(right-left)/2,bottom), fill='#B7D5AF', width=2)
    d.text((left+(right-left)/2+12,top+8),'threshold 0,5',fill='#173B34',font=small)
    img.save(FIG)

doc=Document(); sec=doc.sections[0]
sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(10.5); styles['Normal'].font.color.rgb=RGBColor.from_string(INK)
styles['Body Text'].font.name='Aptos'; styles['Body Text'].font.size=Pt(10.5); styles['Body Text'].paragraph_format.line_spacing=1.12; styles['Body Text'].paragraph_format.space_after=Pt(7)
for level,size in [(1,22),(2,15),(3,11.5)]:
    st=styles[f'Heading {level}']; st.font.name='Aptos Display'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(GREEN); st.paragraph_format.space_before=Pt(12 if level==1 else 8); st.paragraph_format.space_after=Pt(7)
styles['List Bullet'].font.name='Aptos'; styles['List Bullet'].font.size=Pt(10.5)

# Cover
cover=doc.add_table(rows=1, cols=1); cover.alignment=WD_TABLE_ALIGNMENT.CENTER; c=cover.cell(0,0); shade(c,GREEN); set_cell_margins(c,top=500,start=430,bottom=500,end=430)
p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.LEFT; add_run(p,'ML LAB  /  BÁO CÁO HỌC TẬP',bold=True,color='B7D5AF',size=10)
p=c.add_paragraph(style='Title'); p.paragraph_format.space_before=Pt(24); p.paragraph_format.space_after=Pt(6); add_run(p,'Logistic Regression',bold=True,color='FFFFFF',size=34)
p=c.add_paragraph(); p.paragraph_format.space_after=Pt(25); add_run(p,'Từ dữ liệu CSV đến dự đoán nhị phân',color='E7EEE3',size=18,italic=True)
p=c.add_paragraph(); add_run(p,'Một ứng dụng thực hành giúp quan sát toàn bộ quy trình Machine Learning: chuẩn bị dữ liệu, huấn luyện mô hình, đánh giá và dự đoán.',color='FFFFFF',size=11)
doc.add_paragraph()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; add_run(p,'Môn: Machine Learning',color=MUTED,size=10); p.add_run('\n'); add_run(p,'Sinh viên: ........................................',color=MUTED,size=10); p.add_run('\n'); add_run(p,'Ngày nộp: 10/09/2026',color=MUTED,size=10)
doc.add_page_break()

add_heading(doc,'Tóm tắt đề tài',1,'01')
add_body(doc,'ML Lab là một web application nhỏ, server-rendered bằng FastAPI, được xây dựng nhằm minh họa trực quan quy trình Binary Classification với thuật toán Logistic Regression. Người học có thể tải dữ liệu dạng CSV, lựa chọn cột label và các feature, huấn luyện mô hình, đọc các chỉ số đánh giá và thử dự đoán trên một bản ghi mới.')
add_callout(doc,'Mục tiêu học tập','Hiểu được mối liên hệ giữa dữ liệu đầu vào, xác suất dự đoán, ngưỡng phân loại và các metric đánh giá; đồng thời nhận biết giới hạn của kết quả trên dữ liệu nhỏ, dữ liệu giả lập.',PALE,MOSS)
add_heading(doc,'Quy trình năm bước',2)
flow=doc.add_table(rows=2, cols=5); flow.alignment=WD_TABLE_ALIGNMENT.CENTER; flow.autofit=True
steps=[('01','CSV','Mỗi hàng là một ví dụ'),('02','Label + feature','Chọn đáp án và thông tin đầu vào'),('03','Train','Mô hình học trọng số từ dữ liệu train'),('04','Evaluate','Đo đúng/sai trên dữ liệu test'),('05','Predict','Dự đoán một bản ghi mới')]
for i,(n,t,d) in enumerate(steps):
    c=flow.cell(0,i); shade(c,GREEN); set_cell_margins(c,top=130,start=90,bottom=100,end=90); p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; add_run(p,n,bold=True,color='B7D5AF',size=9)
    c=flow.cell(1,i); shade(c,CREAM); set_cell_margins(c,top=130,start=100,bottom=140,end=100); p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; add_run(p,t,bold=True,color=GREEN,size=11); p=c.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; add_run(p,d,color=MUTED,size=8.5)
for row in flow.rows:
    for c in row.cells: set_cell_border(c,top={'val':'single','sz':'4','color':'D5D4CA'},bottom={'val':'single','sz':'4','color':'D5D4CA'},left={'val':'single','sz':'4','color':'D5D4CA'},right={'val':'single','sz':'4','color':'D5D4CA'})
doc.add_page_break()

add_heading(doc,'Cơ sở lý thuyết',1,'02')
add_heading(doc,'Logistic Regression hoạt động như thế nào?',2)
add_body(doc,'Logistic Regression là mô hình dùng cho bài toán phân loại nhị phân. Mô hình kết hợp các feature với những trọng số đã học, sau đó đưa kết quả qua hàm logistic để tạo ra một xác suất trong khoảng từ 0 đến 1.')
make_sigmoid_figure()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(2); p.add_run().add_picture(str(FIG), width=Inches(6.25))
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(8); add_run(p,'Hình 1. Đường cong sigmoid chuyển z thành probability trong khoảng 0 đến 1.',italic=True,color=MUTED,size=8.5)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(10); add_run(p,'p = 1 / (1 + e⁻ᶻ)',bold=True,color=GREEN,size=19)
add_body(doc,'Trong đó, z là tổng có trọng số của các feature. Ví dụ p = 0,83 có thể được hiểu là mô hình ước lượng class dương với xác suất khoảng 83% dựa trên các pattern đã học. Đây là xác suất mô hình, không phải sự chắc chắn hay bằng chứng về quan hệ nhân quả.')
add_callout(doc,'Ngưỡng phân loại','ML Lab sử dụng threshold mặc định bằng 0,5. Nếu p ≥ 0,5, kết quả được hiển thị là class dương; nếu p < 0,5, kết quả là class còn lại.',CREAM,ORANGE)
add_heading(doc,'Thuật ngữ chính',2)
terms=[('Row','Một dòng dữ liệu, đại diện cho một ví dụ.'),('Feature','Thông tin được cung cấp cho mô hình để tạo dự đoán.'),('Label','Đáp án thật mà mô hình cần học.'),('Class dương','Giá trị thứ hai sau khi sắp xếp các label theo thứ tự từ điển.'),('Active model','Phiên bản mô hình đang được sử dụng mặc định ở màn hình Predict.')]
t=doc.add_table(rows=1, cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'; hdr=t.rows[0].cells; hdr[0].text='Thuật ngữ'; hdr[1].text='Giải thích'; set_repeat_table_header(t.rows[0])
for c in hdr: shade(c,GREEN); set_cell_margins(c); 
for c in hdr:
    for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(9.5)
for term,desc in terms:
    cells=t.add_row().cells; cells[0].text=term; cells[1].text=desc
    for i,c in enumerate(cells): set_cell_margins(c); shade(c,PALE if i==0 else 'FFFFFF');
    cells[0].paragraphs[0].runs[0].font.bold=True; cells[0].paragraphs[0].runs[0].font.color.rgb=RGBColor.from_string(GREEN)
doc.add_page_break()

add_heading(doc,'Dữ liệu và quy trình huấn luyện',1,'03')
add_body(doc,'Ứng dụng tiếp nhận CSV mã hóa UTF-8, yêu cầu label có đúng hai giá trị và kiểm tra dữ liệu trước khi train. Sau khi người học chọn label và feature, dữ liệu được chia theo tỷ lệ 80/20 bằng stratified split để giữ tương đối cân bằng hai class trong train và test.')
pipeline=doc.add_table(rows=5, cols=2); pipeline.alignment=WD_TABLE_ALIGNMENT.CENTER; pipeline.style='Table Grid'
pipe_rows=[('Bước 1','Kiểm tra định dạng CSV, schema và giá trị label.'),('Bước 2','Chia train/test theo tỷ lệ 80/20, random_state = 42, stratify theo label.'),('Bước 3','Feature số: điền thiếu bằng median và chuẩn hóa StandardScaler.'),('Bước 4','Feature dạng category: điền giá trị phổ biến và mã hóa OneHotEncoder.'),('Bước 5','Huấn luyện LogisticRegression(max_iter = 2000), sau đó lưu model và metadata.')]
for i,(a,b) in enumerate(pipe_rows):
    pipeline.cell(i,0).text=a; pipeline.cell(i,1).text=b
    for j,c in enumerate(pipeline.rows[i].cells): set_cell_margins(c); shade(c, PALE if j==0 else 'FFFFFF')
    pipeline.cell(i,0).paragraphs[0].runs[0].font.bold=True; pipeline.cell(i,0).paragraphs[0].runs[0].font.color.rgb=RGBColor.from_string(GREEN)
add_callout(doc,'Ngăn ngừa data leakage','Các bước preprocessing chỉ được fit trên tập train. Tập test được giữ lại để đánh giá công bằng khả năng tổng quát hóa của mô hình.',CORAL,ORANGE)
add_heading(doc,'Kết quả được lưu lại',2)
add_bullets(doc,['model.joblib: pipeline đã huấn luyện, dùng cho các lần dự đoán tiếp theo.','metadata.json: tên model, label, feature, class, threshold và thông tin cần thiết cho màn hình Predict.','metrics: các chỉ số đánh giá và confusion matrix của lần train.'])
doc.add_page_break()

add_heading(doc,'Đánh giá mô hình',1,'04')
add_body(doc,'Không nên chỉ nhìn vào Accuracy. Mỗi metric trả lời một câu hỏi khác nhau về kiểu đúng và sai của mô hình. Khi đọc kết quả, cần xem đồng thời metric, confusion matrix và đặc điểm của dataset.')
metrics=[('Accuracy','Trong toàn bộ dự đoán, mô hình đúng bao nhiêu phần?','Có thể gây hiểu nhầm khi một class chiếm đa số.'),('Precision','Trong các lần mô hình đoán class dương, có bao nhiêu lần đúng?','Quan trọng khi báo dương sai gây tốn chi phí.'),('Recall','Trong các class dương thật, mô hình tìm được bao nhiêu?','Quan trọng khi bỏ sót class dương là rủi ro lớn.'),('F1-score','Cân bằng giữa Precision và Recall.','Hữu ích khi hai class không cân bằng.'),('ROC-AUC','Khả năng xếp hạng hai class qua nhiều threshold.','Dựa trên probability, không chỉ nhãn tại 0,5.')]
mt=doc.add_table(rows=1, cols=3); mt.alignment=WD_TABLE_ALIGNMENT.CENTER; mt.style='Table Grid'; heads=['Metric','Câu hỏi metric trả lời','Lưu ý khi diễn giải']
for i,h in enumerate(heads): mt.cell(0,i).text=h; shade(mt.cell(0,i),GREEN); set_cell_margins(mt.cell(0,i));
for c in mt.rows[0].cells:
    for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(9)
set_repeat_table_header(mt.rows[0])
for idx,row in enumerate(metrics):
    cells=mt.add_row().cells
    for i,val in enumerate(row): cells[i].text=val; set_cell_margins(cells[i]); shade(cells[i],PALE if i==0 else ('FFFFFF' if idx%2==0 else 'FBF8F1'))
    cells[0].paragraphs[0].runs[0].font.bold=True; cells[0].paragraphs[0].runs[0].font.color.rgb=RGBColor.from_string(GREEN)
add_heading(doc,'Confusion Matrix',2)
add_body(doc,'Trong confusion matrix, hàng thể hiện label thật và cột thể hiện label dự đoán. Các ô trên đường chéo là dự đoán đúng. Hai loại lỗi còn lại là false positive - báo class dương khi thực tế không phải, và false negative - bỏ sót class dương.')
cm=doc.add_table(rows=3, cols=3); cm.alignment=WD_TABLE_ALIGNMENT.CENTER; cm.style='Table Grid'; vals=[['','Dự đoán âm','Dự đoán dương'],['Thực tế âm','Đúng âm','False positive'],['Thực tế dương','False negative','Đúng dương']]
for i in range(3):
    for j in range(3):
        cm.cell(i,j).text=vals[i][j]; set_cell_margins(cm.cell(i,j),top=150,start=170,bottom=150,end=170); shade(cm.cell(i,j),GREEN if i==0 or j==0 else (CORAL if 'False' in vals[i][j] else PALE))
        if i==0 or j==0:
            for r in cm.cell(i,j).paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
doc.add_page_break()

add_heading(doc,'Ứng dụng minh họa và dữ liệu thực hành',1,'05')
add_body(doc,'ML Lab cung cấp ba bộ dữ liệu CSV hư cấu. Mục đích của các bộ dữ liệu này là giúp người học thực hành chọn feature, quan sát probability và so sánh metric; chúng không đại diện cho cá nhân, khách hàng hay quyết định thực tế.')
samples=[('01','Customer churn','churned: yes / no','days_inactive, tickets, plan','Tăng days_inactive và quan sát probability.'),('02','Marketing response','responded: yes / no','visits, discount_rate, channel, membership','So sánh email và sms ở cùng mức giảm giá.'),('03','Returning customer','returned: yes / no','order_count, average_spend, category, loyalty_tier','Bỏ một feature, train lại và so sánh F1.')]
sg=doc.add_table(rows=1, cols=4); sg.alignment=WD_TABLE_ALIGNMENT.CENTER; sg.style='Table Grid'; sh=['Bộ dữ liệu','Label','Feature','Bài tập gợi ý']
for i,h in enumerate(sh): sg.cell(0,i).text=h; shade(sg.cell(0,i),GREEN); set_cell_margins(sg.cell(0,i));
for c in sg.rows[0].cells:
    for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(9)
set_repeat_table_header(sg.rows[0])
for idx,(n,name,label,features,exercise) in enumerate(samples):
    cells=sg.add_row().cells; vals=[f'{n}. {name}',label,features,exercise]
    for i,val in enumerate(vals): cells[i].text=val; set_cell_margins(cells[i]); shade(cells[i],PALE if i==0 else ('FFFFFF' if idx%2==0 else 'FBF8F1'))
    cells[0].paragraphs[0].runs[0].font.bold=True; cells[0].paragraphs[0].runs[0].font.color.rgb=RGBColor.from_string(GREEN)
add_callout(doc,'Cách sử dụng','Vào Dataset, tải một CSV mẫu hoặc upload file của bạn; chọn label và feature; nhập tên, mục đích model; bấm Train; xem kết quả tại Models và thử một bản ghi mới ở Predict.',PALE,MOSS)
add_heading(doc,'Giới hạn và hướng phát triển',2)
add_bullets(doc,['Dữ liệu mẫu ít và được tạo để minh họa nên score cao không chứng minh mô hình tốt trong thực tế.','Với bài toán thật cần dữ liệu đủ lớn, đại diện, cross-validation và tiêu chí lỗi phù hợp với bối cảnh sử dụng.','Kết quả dự đoán chỉ hỗ trợ học tập; không dùng để tự động đưa ra quyết định ảnh hưởng đến con người.'])
add_heading(doc,'Kết luận',2)
add_body(doc,'Qua ML Lab, người học có thể theo dõi trọn vẹn vòng đời của một bài toán phân loại nhị phân: từ dữ liệu thô đến mô hình, từ xác suất đến nhãn dự đoán, và từ một con số metric đến cách diễn giải có trách nhiệm. Đây là nền tảng để tiếp tục tìm hiểu các mô hình và phương pháp đánh giá nâng cao hơn.')

# Footer and page numbers
for section in doc.sections:
    footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    add_run(footer,'ML Lab  •  Báo cáo Logistic Regression  |  ',color=MUTED,size=8)
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)

doc.save(path)
print(path)
